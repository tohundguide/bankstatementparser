"""
Excel Exporter
==============
Exports parsed bank statement data to professionally formatted Excel files.

Features:
  - Beautifully formatted transaction sheet
  - Monthly summary sheet with totals & charts
  - Auto-filter and freeze panes
  - Color-coded withdrawals (red) and deposits (green)
  - CSV export option
"""

import csv
import os
import re
from collections import defaultdict
from datetime import datetime
from typing import Dict, List

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill, numbers
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, Reference


def export_to_excel(result: Dict, output_path: str):
    """
    Export parsed transaction data to a formatted Excel file with:
      - Sheet 1: Full transaction list (professional formatting)
      - Sheet 2: Monthly summary with totals
      - Sheet 3: Narration mapping (unique particulars for user classification)
    
    Args:
        result: Parsed result dict from a bank parser
        output_path: Path to save the Excel file
    """
    wb = Workbook()
    
    # ── Create Transaction Sheet ──
    _create_transaction_sheet(wb, result)
    
    # ── Create Monthly Summary Sheet ──
    _create_summary_sheet(wb, result)
    
    # ── Create Narration Mapping Sheet ──
    _create_narration_sheet(wb, result)
    
    # ── Save ──
    wb.save(output_path)
    print(f"  [OK] Excel saved: {output_path}")


def export_to_csv(result: Dict, output_path: str):
    """
    Export parsed transaction data to a CSV file.
    
    Args:
        result: Parsed result dict from a bank parser
        output_path: Path to save the CSV file
    """
    transactions = result.get('transactions', [])
    
    with open(output_path, 'w', newline='', encoding='utf-8-sig') as f:
        writer = csv.writer(f)
        
        # Header info
        writer.writerow([f"Bank: {result.get('bank_name', 'N/A')}"])
        account_info = result.get('account_info', {})
        if account_info.get('account_number'):
            writer.writerow([f"Account: {account_info['account_number']}"])
        if account_info.get('account_holder'):
            writer.writerow([f"Holder: {account_info['account_holder']}"])
        writer.writerow([f"Period: {result.get('period', 'N/A')}"])
        writer.writerow([f"Total Transactions: {len(transactions)}"])
        writer.writerow([])
        
        # Column headers
        writer.writerow(['Date', 'Particulars', 'Chq/Ref No.', 'Withdrawals', 'Deposits', 'Balance', 'Verification'])
        
        # Data rows
        total_w = 0
        total_d = 0
        prev_bal = None
        csv_match = 0
        csv_check = 0
        for idx, txn in enumerate(transactions):
            bal_str = txn.get('balance', '')
            w_str = txn.get('withdrawal', '')
            d_str = txn.get('deposit', '')
            parsed_bal = _parse_balance_value(bal_str)
            
            # Compute verification
            if idx == 0:
                chk = '○ Opening'
            elif prev_bal is not None and parsed_bal is not None:
                try:
                    wv = float(str(w_str).replace(',', '')) if w_str else 0
                    dv = float(str(d_str).replace(',', '')) if d_str else 0
                except (ValueError, TypeError):
                    wv, dv = 0, 0
                exp = prev_bal + dv - wv
                diff = parsed_bal - exp
                csv_check += 1
                if abs(diff) < 0.02:
                    chk = '✓ Match'
                    csv_match += 1
                else:
                    chk = f'✗ Diff: {diff:+,.2f}'
            else:
                chk = '? N/A'
            prev_bal = parsed_bal
            
            writer.writerow([
                txn.get('date', ''),
                txn.get('particulars', ''),
                txn.get('chq_ref', ''),
                w_str,
                d_str,
                bal_str,
                chk,
            ])
            try:
                if w_str:
                    total_w += float(str(w_str).replace(',', ''))
            except (ValueError, TypeError):
                pass
            try:
                if d_str:
                    total_d += float(str(d_str).replace(',', ''))
            except (ValueError, TypeError):
                pass
        
        # Totals
        writer.writerow([])
        csv_pct = (csv_match / csv_check * 100) if csv_check > 0 else 0
        writer.writerow(['', '', 'TOTALS', f'{total_w:.2f}', f'{total_d:.2f}', '', f'Verified: {csv_match}/{csv_check} ({csv_pct:.0f}%)'])
    
    print(f"  [OK] CSV saved: {output_path}")


# ─── INTERNAL HELPERS ───


def _parse_balance_value(balance_str):
    """Parse balance string to signed float. Cr/positive = +, Dr/negative = -."""
    if not balance_str:
        return None
    s = str(balance_str).strip()
    is_dr = False
    # Check for Dr/DR suffix
    for suffix in ['Dr.', 'DR.', 'Dr', 'DR']:
        if s.endswith(suffix):
            is_dr = True
            s = s[:-len(suffix)].strip()
            break
    else:
        # Check for Cr/CR suffix (positive, just strip)
        for suffix in ['Cr.', 'CR.', 'Cr', 'CR']:
            if s.endswith(suffix):
                s = s[:-len(suffix)].strip()
                break
    s = s.replace(',', '').strip()
    try:
        val = float(s)
        return -val if is_dr else val
    except (ValueError, TypeError):
        return None


def _create_transaction_sheet(wb: Workbook, result: Dict):
    """Create the main transactions sheet with professional formatting."""
    ws = wb.active
    ws.title = "Transactions"
    
    # ── Styles ──
    header_font = Font(name='Calibri', bold=True, size=14, color='FFFFFF')
    subheader_font = Font(name='Calibri', bold=True, size=10, color='333333')
    col_header_font = Font(name='Calibri', bold=True, size=11, color='FFFFFF')
    data_font = Font(name='Calibri', size=10)
    
    header_fill = PatternFill(start_color='1B4F72', end_color='1B4F72', fill_type='solid')
    col_header_fill = PatternFill(start_color='2E86C1', end_color='2E86C1', fill_type='solid')
    alt_row_fill = PatternFill(start_color='EBF5FB', end_color='EBF5FB', fill_type='solid')
    withdrawal_fill = PatternFill(start_color='FDEDEC', end_color='FDEDEC', fill_type='solid')
    deposit_fill = PatternFill(start_color='EAFAF1', end_color='EAFAF1', fill_type='solid')
    totals_fill = PatternFill(start_color='D5F5E3', end_color='D5F5E3', fill_type='solid')
    
    thin_border = Border(
        left=Side(style='thin', color='BDC3C7'),
        right=Side(style='thin', color='BDC3C7'),
        top=Side(style='thin', color='BDC3C7'),
        bottom=Side(style='thin', color='BDC3C7'),
    )
    
    center_align = Alignment(horizontal='center', vertical='center', wrap_text=True)
    left_align = Alignment(horizontal='left', vertical='center', wrap_text=True)
    right_align = Alignment(horizontal='right', vertical='center')
    
    # ── Title ──
    row = 1
    ws.merge_cells(f'A{row}:H{row}')
    title_cell = ws.cell(row=row, column=1, value=f"{result['bank_name']} — Account Statement")
    title_cell.font = header_font
    title_cell.fill = header_fill
    title_cell.alignment = center_align
    ws.row_dimensions[row].height = 35
    
    # ── Account Info ──
    account_info = result.get('account_info', {})
    row = 2
    info_items = [
        ('Account No:', account_info.get('account_number', 'N/A')),
        ('Account Holder:', account_info.get('account_holder', 'N/A')),
        ('Account Type:', account_info.get('type', 'N/A')),
        ('IFSC:', account_info.get('ifsc', 'N/A')),
        ('Period:', result.get('period', 'N/A')),
        ('Total Transactions:', str(len(result.get('transactions', [])))),
    ]
    
    for label, value in info_items:
        ws.cell(row=row, column=1, value=label).font = Font(name='Calibri', bold=True, size=9, color='555555')
        ws.cell(row=row, column=2, value=value).font = Font(name='Calibri', size=9, color='333333')
        ws.merge_cells(f'B{row}:H{row}')
        row += 1
    
    row += 1  # Empty row
    
    # ── Column Headers ──
    headers = ['DATE', 'PARTICULARS', 'CHQ.NO./REF.NO.', 'WITHDRAWALS', 'DEPOSITS', 'BALANCE', 'VERIFICATION', 'NARRATION']
    col_widths = [14, 50, 22, 16, 16, 20, 18, 40]
    
    for col_idx, (header, width) in enumerate(zip(headers, col_widths), 1):
        cell = ws.cell(row=row, column=col_idx, value=header)
        cell.font = col_header_font
        cell.fill = col_header_fill
        cell.alignment = center_align
        cell.border = thin_border
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    
    ws.row_dimensions[row].height = 25
    header_row = row
    row += 1
    
    # ── Transaction Data ──
    transactions = result.get('transactions', [])
    
    # Balance verification tracking
    prev_parsed_balance = None
    match_count = 0
    check_count = 0
    
    for idx, txn in enumerate(transactions):
        # Date
        ws.cell(row=row, column=1, value=txn['date']).font = data_font
        ws.cell(row=row, column=1).alignment = center_align
        
        # Particulars
        ws.cell(row=row, column=2, value=txn['particulars']).font = data_font
        ws.cell(row=row, column=2).alignment = left_align
        
        # Cheque/Ref
        ws.cell(row=row, column=3, value=txn['chq_ref']).font = data_font
        ws.cell(row=row, column=3).alignment = center_align
        
        # Withdrawal
        withdrawal_val = txn['withdrawal']
        if withdrawal_val:
            try:
                ws.cell(row=row, column=4, value=float(str(withdrawal_val).replace(',', ''))).font = Font(name='Calibri', size=10, color='C0392B')
                ws.cell(row=row, column=4).number_format = '#,##0.00'
            except ValueError:
                ws.cell(row=row, column=4, value=withdrawal_val).font = data_font
        ws.cell(row=row, column=4).alignment = right_align
        
        # Deposit
        deposit_val = txn['deposit']
        if deposit_val:
            try:
                ws.cell(row=row, column=5, value=float(str(deposit_val).replace(',', ''))).font = Font(name='Calibri', size=10, color='27AE60')
                ws.cell(row=row, column=5).number_format = '#,##0.00'
            except ValueError:
                ws.cell(row=row, column=5, value=deposit_val).font = data_font
        ws.cell(row=row, column=5).alignment = right_align
        
        # Balance
        balance_val = txn['balance']
        ws.cell(row=row, column=6, value=balance_val).font = data_font
        ws.cell(row=row, column=6).alignment = right_align
        
        # Verification (column 7)
        parsed_bal = _parse_balance_value(balance_val)
        if idx == 0:
            check_text = "○ Opening"
            check_color = '7F8C8D'
        elif prev_parsed_balance is not None and parsed_bal is not None:
            try:
                w_val = float(str(withdrawal_val).replace(',', '')) if withdrawal_val else 0
                d_val = float(str(deposit_val).replace(',', '')) if deposit_val else 0
            except (ValueError, TypeError):
                w_val, d_val = 0, 0
            expected = prev_parsed_balance + d_val - w_val
            diff = parsed_bal - expected
            check_count += 1
            if abs(diff) < 0.02:
                check_text = "✓ Match"
                check_color = '27AE60'
                match_count += 1
            else:
                check_text = f"✗ Diff: {diff:+,.2f}"
                check_color = 'C0392B'
        else:
            check_text = "? N/A"
            check_color = '7F8C8D'
        
        check_cell = ws.cell(row=row, column=7, value=check_text)
        check_cell.font = Font(name='Calibri', size=9, color=check_color, bold=check_text.startswith('✗'))
        check_cell.alignment = center_align
        prev_parsed_balance = parsed_bal
        
        # Narration (column 8) — VLOOKUP from Narration Mapping sheet
        vlookup = f"=IFERROR(VLOOKUP(B{row},'Narration Mapping'!A:B,2,FALSE),\"\")"
        narr_cell = ws.cell(row=row, column=8, value=vlookup)
        narr_cell.font = Font(name='Calibri', size=10, color='2C3E50', italic=True)
        narr_cell.alignment = left_align
        
        # Row styling
        row_fill = alt_row_fill if idx % 2 == 0 else None
        if withdrawal_val and not deposit_val:
            row_fill = withdrawal_fill
        elif deposit_val and not withdrawal_val:
            row_fill = deposit_fill
        
        for col_idx in range(1, 9):
            cell = ws.cell(row=row, column=col_idx)
            cell.border = thin_border
            if row_fill:
                cell.fill = row_fill
        
        row += 1
    
    # ── Totals Row ──
    row += 1
    ws.merge_cells(f'A{row}:C{row}')
    summary_cell = ws.cell(row=row, column=1, value='TOTALS')
    summary_cell.font = Font(name='Calibri', bold=True, size=11)
    summary_cell.alignment = center_align
    summary_cell.fill = totals_fill
    
    total_withdrawals = 0
    total_deposits = 0
    for txn in transactions:
        try:
            if txn['withdrawal']:
                total_withdrawals += float(str(txn['withdrawal']).replace(',', ''))
        except (ValueError, TypeError):
            pass
        try:
            if txn['deposit']:
                total_deposits += float(str(txn['deposit']).replace(',', ''))
        except (ValueError, TypeError):
            pass
    
    wd_cell = ws.cell(row=row, column=4, value=total_withdrawals)
    wd_cell.font = Font(name='Calibri', bold=True, size=11, color='C0392B')
    wd_cell.number_format = '#,##0.00'
    wd_cell.alignment = right_align
    wd_cell.border = thin_border
    
    dp_cell = ws.cell(row=row, column=5, value=total_deposits)
    dp_cell.font = Font(name='Calibri', bold=True, size=11, color='27AE60')
    dp_cell.number_format = '#,##0.00'
    dp_cell.alignment = right_align
    dp_cell.border = thin_border
    
    # Verification summary in column 7
    verify_pct = (match_count / check_count * 100) if check_count > 0 else 0
    verify_text = f"Verified: {match_count}/{check_count} ({verify_pct:.0f}%)"
    verify_cell = ws.cell(row=row, column=7, value=verify_text)
    verify_cell.font = Font(name='Calibri', bold=True, size=9, color='27AE60' if verify_pct == 100 else 'E67E22')
    verify_cell.alignment = center_align
    
    for col_idx in range(1, 9):
        ws.cell(row=row, column=col_idx).border = thin_border
        ws.cell(row=row, column=col_idx).fill = totals_fill
    
    # ── Net Flow ──
    row += 1
    ws.merge_cells(f'A{row}:C{row}')
    net_label = ws.cell(row=row, column=1, value='NET FLOW (Deposits - Withdrawals)')
    net_label.font = Font(name='Calibri', bold=True, size=10, color='555555')
    net_label.alignment = center_align
    
    net = total_deposits - total_withdrawals
    ws.merge_cells(f'D{row}:E{row}')
    net_cell = ws.cell(row=row, column=4, value=net)
    net_cell.font = Font(name='Calibri', bold=True, size=12, color='27AE60' if net >= 0 else 'C0392B')
    net_cell.number_format = '#,##0.00'
    net_cell.alignment = right_align
    
    # ── Freeze + Filter ──
    ws.freeze_panes = f'A{header_row + 1}'
    ws.auto_filter.ref = f'A{header_row}:H{header_row + len(transactions)}'


def _create_summary_sheet(wb: Workbook, result: Dict):
    """Create a monthly summary sheet with totals and chart."""
    ws = wb.create_sheet(title="Monthly Summary")
    
    transactions = result.get('transactions', [])
    if not transactions:
        ws.cell(row=1, column=1, value="No transactions to summarize")
        return
    
    # ── Parse dates and aggregate by month ──
    monthly = defaultdict(lambda: {'withdrawals': 0, 'deposits': 0, 'count': 0})
    
    date_formats = [
        '%d-%m-%Y', '%d/%m/%Y', '%d-%b-%Y', '%d/%b/%Y',
        '%d-%m-%y', '%d/%m/%y', '%d-%b-%y', '%d/%b/%y',
        '%Y-%m-%d', '%m/%d/%Y',
    ]
    
    for txn in transactions:
        date_str = txn.get('date', '').strip()
        if not date_str:
            continue
        
        parsed_date = None
        for fmt in date_formats:
            try:
                parsed_date = datetime.strptime(date_str, fmt)
                break
            except ValueError:
                continue
        
        if not parsed_date:
            # Try extracting just digits-based date
            m = re.match(r'(\d{1,2})[/\-](\d{1,2})[/\-](\d{2,4})', date_str)
            if m:
                day, month, year = m.groups()
                if len(year) == 2:
                    year = '20' + year
                try:
                    parsed_date = datetime(int(year), int(month), int(day))
                except (ValueError, TypeError):
                    continue
            else:
                continue
        
        month_key = parsed_date.strftime('%Y-%m')
        month_label = parsed_date.strftime('%b %Y')
        
        monthly[month_key]['label'] = month_label
        monthly[month_key]['count'] += 1
        
        try:
            w = txn.get('withdrawal', '')
            if w:
                monthly[month_key]['withdrawals'] += float(str(w).replace(',', ''))
        except (ValueError, TypeError):
            pass
        
        try:
            d = txn.get('deposit', '')
            if d:
                monthly[month_key]['deposits'] += float(str(d).replace(',', ''))
        except (ValueError, TypeError):
            pass
    
    if not monthly:
        ws.cell(row=1, column=1, value="Could not parse dates for monthly summary")
        return
    
    # Sort by month
    sorted_months = sorted(monthly.items())
    
    # ── Styles ──
    header_fill = PatternFill(start_color='1B4F72', end_color='1B4F72', fill_type='solid')
    header_font = Font(name='Calibri', bold=True, size=12, color='FFFFFF')
    col_header_fill = PatternFill(start_color='2E86C1', end_color='2E86C1', fill_type='solid')
    col_header_font = Font(name='Calibri', bold=True, size=11, color='FFFFFF')
    data_font = Font(name='Calibri', size=10)
    totals_fill = PatternFill(start_color='D5F5E3', end_color='D5F5E3', fill_type='solid')
    alt_fill = PatternFill(start_color='EBF5FB', end_color='EBF5FB', fill_type='solid')
    
    thin_border = Border(
        left=Side(style='thin', color='BDC3C7'),
        right=Side(style='thin', color='BDC3C7'),
        top=Side(style='thin', color='BDC3C7'),
        bottom=Side(style='thin', color='BDC3C7'),
    )
    center_align = Alignment(horizontal='center', vertical='center')
    right_align = Alignment(horizontal='right', vertical='center')
    
    # ── Title ──
    row = 1
    ws.merge_cells(f'A{row}:F{row}')
    title = ws.cell(row=row, column=1, value=f"{result['bank_name']} — Monthly Summary")
    title.font = header_font
    title.fill = header_fill
    title.alignment = center_align
    ws.row_dimensions[row].height = 30
    
    row = 3
    # ── Column Headers ──
    headers = ['Month', 'Transactions', 'Total Withdrawals', 'Total Deposits', 'Net Flow', 'Avg. Transaction']
    widths = [15, 14, 20, 20, 18, 18]
    
    for col_idx, (h, w) in enumerate(zip(headers, widths), 1):
        cell = ws.cell(row=row, column=col_idx, value=h)
        cell.font = col_header_font
        cell.fill = col_header_fill
        cell.alignment = center_align
        cell.border = thin_border
        ws.column_dimensions[get_column_letter(col_idx)].width = w
    
    summary_header_row = row
    row += 1
    
    # ── Monthly Data ──
    grand_w = 0
    grand_d = 0
    grand_count = 0
    
    for idx, (key, data) in enumerate(sorted_months):
        net = data['deposits'] - data['withdrawals']
        avg_txn = (data['withdrawals'] + data['deposits']) / data['count'] if data['count'] > 0 else 0
        
        ws.cell(row=row, column=1, value=data.get('label', key)).font = Font(name='Calibri', bold=True, size=10)
        ws.cell(row=row, column=1).alignment = center_align
        
        ws.cell(row=row, column=2, value=data['count']).font = data_font
        ws.cell(row=row, column=2).alignment = center_align
        
        ws.cell(row=row, column=3, value=data['withdrawals']).font = Font(name='Calibri', size=10, color='C0392B')
        ws.cell(row=row, column=3).number_format = '#,##0.00'
        ws.cell(row=row, column=3).alignment = right_align
        
        ws.cell(row=row, column=4, value=data['deposits']).font = Font(name='Calibri', size=10, color='27AE60')
        ws.cell(row=row, column=4).number_format = '#,##0.00'
        ws.cell(row=row, column=4).alignment = right_align
        
        ws.cell(row=row, column=5, value=net).font = Font(name='Calibri', size=10, color='27AE60' if net >= 0 else 'C0392B')
        ws.cell(row=row, column=5).number_format = '#,##0.00'
        ws.cell(row=row, column=5).alignment = right_align
        
        ws.cell(row=row, column=6, value=avg_txn).font = data_font
        ws.cell(row=row, column=6).number_format = '#,##0.00'
        ws.cell(row=row, column=6).alignment = right_align
        
        # Alternate row fill
        if idx % 2 == 0:
            for c in range(1, 7):
                ws.cell(row=row, column=c).fill = alt_fill
        
        for c in range(1, 7):
            ws.cell(row=row, column=c).border = thin_border
        
        grand_w += data['withdrawals']
        grand_d += data['deposits']
        grand_count += data['count']
        row += 1
    
    # ── Grand Totals ──
    row += 1
    ws.cell(row=row, column=1, value='GRAND TOTAL').font = Font(name='Calibri', bold=True, size=11)
    ws.cell(row=row, column=1).alignment = center_align
    ws.cell(row=row, column=2, value=grand_count).font = Font(name='Calibri', bold=True, size=11)
    ws.cell(row=row, column=2).alignment = center_align
    
    ws.cell(row=row, column=3, value=grand_w)
    ws.cell(row=row, column=3).font = Font(name='Calibri', bold=True, size=11, color='C0392B')
    ws.cell(row=row, column=3).number_format = '#,##0.00'
    ws.cell(row=row, column=3).alignment = right_align
    
    ws.cell(row=row, column=4, value=grand_d)
    ws.cell(row=row, column=4).font = Font(name='Calibri', bold=True, size=11, color='27AE60')
    ws.cell(row=row, column=4).number_format = '#,##0.00'
    ws.cell(row=row, column=4).alignment = right_align
    
    grand_net = grand_d - grand_w
    ws.cell(row=row, column=5, value=grand_net)
    ws.cell(row=row, column=5).font = Font(name='Calibri', bold=True, size=11, color='27AE60' if grand_net >= 0 else 'C0392B')
    ws.cell(row=row, column=5).number_format = '#,##0.00'
    ws.cell(row=row, column=5).alignment = right_align
    
    for c in range(1, 7):
        ws.cell(row=row, column=c).fill = totals_fill
        ws.cell(row=row, column=c).border = thin_border
    
    # ── Bar Chart ──
    if len(sorted_months) >= 2:
        try:
            chart = BarChart()
            chart.title = "Monthly Cash Flow"
            chart.y_axis.title = "Amount (₹)"
            chart.x_axis.title = "Month"
            chart.style = 10
            chart.width = 25
            chart.height = 14
            
            data_start = summary_header_row + 1
            data_end = data_start + len(sorted_months) - 1
            
            # Withdrawals series
            wd_ref = Reference(ws, min_col=3, min_row=summary_header_row, max_row=data_end)
            chart.add_data(wd_ref, titles_from_data=True)
            
            # Deposits series
            dp_ref = Reference(ws, min_col=4, min_row=summary_header_row, max_row=data_end)
            chart.add_data(dp_ref, titles_from_data=True)
            
            # Categories (month labels)
            cats = Reference(ws, min_col=1, min_row=data_start, max_row=data_end)
            chart.set_categories(cats)
            
            # Color the series
            chart.series[0].graphicalProperties.solidFill = "E74C3C"  # Red for withdrawals
            chart.series[1].graphicalProperties.solidFill = "2ECC71"  # Green for deposits
            
            chart_row = row + 3
            ws.add_chart(chart, f"A{chart_row}")
        except Exception:
            pass  # Chart creation is optional


def _create_narration_sheet(wb: Workbook, result: Dict):
    """Create a narration mapping sheet with unique particulars for user classification."""
    ws = wb.create_sheet(title="Narration Mapping")
    
    transactions = result.get('transactions', [])
    if not transactions:
        ws.cell(row=1, column=1, value="No transactions to map")
        return
    
    # ── Extract unique particulars (preserving first-occurrence order) ──
    seen = set()
    unique_particulars = []
    for txn in transactions:
        p = txn.get('particulars', '').strip()
        if p and p not in seen:
            seen.add(p)
            unique_particulars.append(p)
    
    if not unique_particulars:
        ws.cell(row=1, column=1, value="No particulars found")
        return
    
    # ── Styles ──
    header_fill = PatternFill(start_color='1B4F72', end_color='1B4F72', fill_type='solid')
    header_font = Font(name='Calibri', bold=True, size=12, color='FFFFFF')
    col_header_fill = PatternFill(start_color='2E86C1', end_color='2E86C1', fill_type='solid')
    col_header_font = Font(name='Calibri', bold=True, size=11, color='FFFFFF')
    data_font = Font(name='Calibri', size=10)
    input_fill = PatternFill(start_color='FFF9C4', end_color='FFF9C4', fill_type='solid')  # Light yellow
    alt_fill = PatternFill(start_color='EBF5FB', end_color='EBF5FB', fill_type='solid')
    
    thin_border = Border(
        left=Side(style='thin', color='BDC3C7'),
        right=Side(style='thin', color='BDC3C7'),
        top=Side(style='thin', color='BDC3C7'),
        bottom=Side(style='thin', color='BDC3C7'),
    )
    center_align = Alignment(horizontal='center', vertical='center', wrap_text=True)
    left_align = Alignment(horizontal='left', vertical='center', wrap_text=True)
    
    # ── Title ──
    row = 1
    ws.merge_cells(f'A{row}:B{row}')
    title = ws.cell(row=row, column=1, value="Narration Mapping — Type your classifications in Column B")
    title.font = header_font
    title.fill = header_fill
    title.alignment = center_align
    ws.row_dimensions[row].height = 35
    
    # ── Instructions ──
    row = 2
    ws.merge_cells(f'A{row}:B{row}')
    instr = ws.cell(row=row, column=1,
                    value="Fill Column B with detailed narrations. They will auto-populate in the Transactions sheet Column H via VLOOKUP.")
    instr.font = Font(name='Calibri', size=9, italic=True, color='555555')
    instr.alignment = center_align
    ws.row_dimensions[row].height = 20
    
    row = 3
    # ── Column Headers ──
    ws.column_dimensions['A'].width = 55
    ws.column_dimensions['B'].width = 50
    
    h1 = ws.cell(row=row, column=1, value='PARTICULARS (from statement)')
    h1.font = col_header_font
    h1.fill = col_header_fill
    h1.alignment = center_align
    h1.border = thin_border
    
    h2 = ws.cell(row=row, column=2, value='DETAILED NARRATION (type here)')
    h2.font = col_header_font
    h2.fill = col_header_fill
    h2.alignment = center_align
    h2.border = thin_border
    
    ws.row_dimensions[row].height = 25
    row += 1
    
    # ── Data Rows ──
    for idx, particular in enumerate(unique_particulars):
        # Column A: Particulars (locked reference)
        a_cell = ws.cell(row=row, column=1, value=particular)
        a_cell.font = data_font
        a_cell.alignment = left_align
        a_cell.border = thin_border
        if idx % 2 == 0:
            a_cell.fill = alt_fill
        
        # Column B: Empty for user input (yellow highlight)
        b_cell = ws.cell(row=row, column=2, value='')
        b_cell.font = Font(name='Calibri', size=10, color='2C3E50')
        b_cell.alignment = left_align
        b_cell.border = thin_border
        b_cell.fill = input_fill
        
        row += 1
    
    # ── Summary ──
    row += 1
    ws.cell(row=row, column=1, value=f"Total unique particulars: {len(unique_particulars)}").font = Font(
        name='Calibri', size=9, italic=True, color='7F8C8D')
    
    # Freeze header
    ws.freeze_panes = 'A4'

