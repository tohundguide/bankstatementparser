"""
Text Extractor Module
=====================
Extracts raw text from various file formats: PDF, DOCX, TXT, CSV.
Supports:
  - Regular PDFs (text-based) via pdfplumber / PyPDF2
  - Scanned/image PDFs via OCR (Tesseract – free, open-source)
  - Password-protected PDFs
  - DOCX files
  - Plain text / CSV files
"""

import os
import re
import logging
import shutil
import threading
from concurrent.futures import ThreadPoolExecutor

# Limit Tesseract to 1 OpenMP thread per job.
# Must be set BEFORE importing pytesseract or calling tesseract.
# Without this, Tesseract uses ALL CPU cores — 2 concurrent jobs = 100% CPU saturation.
os.environ['OMP_THREAD_LIMIT'] = '1'
os.environ['OMP_NUM_THREADS'] = '1'

logger = logging.getLogger(__name__)


def extract_text(filepath: str, ext: str, password: str = None) -> str:
    """
    Extract raw text from a file based on its extension.
    
    Args:
        filepath: Absolute path to the file
        ext: File extension (lowercase, with dot, e.g. '.pdf')
        password: Optional password for encrypted PDFs
    
    Returns:
        Raw text content of the file
    """
    ext = ext.lower()
    
    if ext in ['.txt', '.csv']:
        return _extract_from_text(filepath)
    elif ext in ['.docx', '.doc']:
        return _extract_from_docx(filepath)
    elif ext == '.pdf':
        return _extract_from_pdf(filepath, password=password)
    else:
        raise ValueError(f"Unsupported file format: {ext}")


def _extract_from_text(filepath: str) -> str:
    """Extract text from plain text or CSV files."""
    encodings = ['utf-8', 'utf-8-sig', 'latin-1', 'cp1252', 'ascii']
    
    for encoding in encodings:
        try:
            with open(filepath, 'r', encoding=encoding) as f:
                return f.read()
        except (UnicodeDecodeError, UnicodeError):
            continue
    
    # Fallback: read as binary and decode with errors='replace'
    with open(filepath, 'rb') as f:
        return f.read().decode('utf-8', errors='replace')


def _extract_from_docx(filepath: str) -> str:
    """Extract text from DOCX files, preserving line structure."""
    try:
        from docx import Document
    except ImportError:
        raise ImportError("python-docx is required for DOCX files. Install: pip install python-docx")
    
    doc = Document(filepath)
    lines = []
    for para in doc.paragraphs:
        lines.append(para.text)
    
    # Also extract text from tables (bank statements often use tables in DOCX)
    for table in doc.tables:
        for row in table.rows:
            row_text = []
            for cell in row.cells:
                cell_text = cell.text.strip()
                if cell_text:
                    row_text.append(cell_text)
            if row_text:
                lines.append('  '.join(row_text))
    
    return '\n'.join(lines)


def _extract_from_pdf(filepath: str, password: str = None) -> str:
    """
    Extract text from PDF files.
    Pipeline:
      1. Try pdfplumber (best for structured/table PDFs)
      2. Try PyPDF2 as fallback
      3. If both fail (scanned PDF), try OCR with Tesseract
    
    Handles password-protected PDFs automatically.
    """
    text = ""
    
    # --- Step 1: Try pdfplumber ---
    try:
        import pdfplumber
        open_kwargs = {}
        if password:
            open_kwargs['password'] = password
        
        with pdfplumber.open(filepath, **open_kwargs) as pdf:
            pages = []
            for page in pdf.pages:
                page_text = page.extract_text()
                if page_text:
                    pages.append(page_text)
                # Release this page's parsed objects/layout cache before moving on.
                # Without this, pdfplumber retains every page's objects in RAM, so a
                # large multi-hundred-page PDF grows memory linearly and can OOM.
                # With it, peak memory stays roughly flat regardless of page count.
                try:
                    page.flush_cache()
                except Exception:
                    pass
            text = '\n'.join(pages)
            if text.strip():
                logger.info(f"Extracted {len(text)} chars via pdfplumber")
                return text
    except Exception as e:
        logger.debug(f"pdfplumber failed: {e}")
    
    # --- Step 2: Try PyPDF2 ---
    try:
        from PyPDF2 import PdfReader
        reader = PdfReader(filepath)
        
        # Handle encrypted PDFs
        if reader.is_encrypted:
            if password:
                reader.decrypt(password)
            else:
                # Try common bank statement passwords
                common_passwords = _get_common_passwords(filepath)
                decrypted = False
                for pwd in common_passwords:
                    try:
                        if reader.decrypt(pwd):
                            decrypted = True
                            logger.info(f"PDF decrypted with common password pattern")
                            break
                    except Exception:
                        continue
                if not decrypted:
                    raise ValueError(
                        "This PDF is password-protected. Please enter the password. "
                        "Common passwords for bank statements: your account number, DOB (DDMMYYYY), or PAN number."
                    )
        
        pages = []
        for page in reader.pages:
            page_text = page.extract_text()
            if page_text:
                pages.append(page_text)
        text = '\n'.join(pages)
        if text.strip():
            logger.info(f"Extracted {len(text)} chars via PyPDF2")
            return text
    except ValueError:
        raise  # Re-raise password errors
    except Exception as e:
        logger.debug(f"PyPDF2 failed: {e}")
    
    # --- Step 3: OCR fallback for scanned PDFs ---
    if not text.strip():
        logger.info("No text extracted — attempting OCR...")
        try:
            text = _extract_via_ocr(filepath)
            if text.strip():
                logger.info(f"Extracted {len(text)} chars via OCR")
                return text
        except ImportError as e:
            raise ValueError(
                f"This appears to be a scanned/image-based PDF. "
                f"OCR is required but dependencies are missing: {str(e)}. "
                f"Install: pip install pytesseract pdf2image Pillow "
                f"and install Tesseract from https://github.com/UB-Mannheim/tesseract/wiki"
            )
        except Exception as e:
            raise ValueError(
                f"Could not extract text from PDF (scanned/image-based). "
                f"OCR failed: {str(e)}. "
                f"Ensure Tesseract is installed and in your PATH."
            )
    
    raise ValueError("Could not extract text from PDF. The file may be empty or corrupted.")


def _get_common_passwords(filepath: str) -> list:
    """
    Generate common bank statement passwords to try.
    Banks usually use: account number, DOB, PAN, etc.
    Returns a list of passwords to attempt.
    """
    passwords = ['', ' ']
    
    # Extract potential account numbers from filename
    filename = os.path.basename(filepath)
    # Find sequences of digits in filename
    digit_sequences = re.findall(r'\d{4,}', filename)
    passwords.extend(digit_sequences)
    
    return passwords


def _extract_via_ocr(filepath: str) -> str:
    """
    Extract text from a scanned / text-less PDF with Tesseract.

    Pages are rasterised at OCR_DPI (default 150 — enough for statement text;
    300 quadruples the CPU and, measured on outline-font statements, does not
    read digits any better) and recognised OCR_WORKERS pages at a time.
    Tesseract itself stays single-threaded (OMP_THREAD_LIMIT=1 above), so the
    parallelism is one process per page and CPU use is bounded by OCR_WORKERS.

    Rendering uses poppler (pdf2image) when it is installed — the Docker image
    ships it — and otherwise pypdfium2, which pdfplumber already depends on.
    Either way a page is rendered only when its turn comes, so memory stays
    flat for long scans instead of holding every page image at once.

    Requirements (all free):
      - pytesseract + the Tesseract binary (apt: tesseract-ocr / UB-Mannheim installer)
      - pdf2image + poppler (apt: poppler-utils)  OR  pypdfium2 (pip)
    """
    try:
        import pytesseract
    except ImportError:
        raise ImportError(
            "pytesseract is required for OCR. Install: pip install pytesseract"
        )
    _configure_tesseract_cmd(pytesseract)

    poppler_path = _find_poppler_path()
    use_poppler = poppler_path is not None
    if not use_poppler and not _pdfium_available():
        raise RuntimeError(
            "No PDF rasteriser available: install poppler (Linux: apt install "
            "poppler-utils; Windows: https://github.com/oschwartz10612/poppler-windows/releases "
            "extracted to C:\\poppler) or `pip install pypdfium2`."
        )

    page_count = _page_count(filepath, poppler_path)
    dpi = int(os.environ.get('OCR_DPI', '150'))
    workers = max(1, min(int(os.environ.get('OCR_WORKERS', '3')), os.cpu_count() or 1, page_count))
    logger.info(f"OCR: {page_count} page(s) at {dpi} dpi, {workers} worker(s), "
                f"renderer={'poppler' if use_poppler else 'pdfium'}")

    def ocr_page(page_no: int) -> str:
        try:
            if use_poppler:
                try:
                    img = _render_page_poppler(filepath, page_no, dpi, poppler_path)
                except Exception as e:
                    if not _pdfium_available():
                        raise
                    logger.warning(f"  poppler failed on page {page_no} ({e}); using pdfium")
                    img = _render_page_pdfium(filepath, page_no, dpi)
            else:
                img = _render_page_pdfium(filepath, page_no, dpi)
        except Exception as e:
            raise RuntimeError(f"Could not convert PDF page {page_no} to an image: {e}")
        try:
            text = pytesseract.image_to_string(img, config=_OCR_CONFIG)
        except Exception as e:
            logger.warning(f"  OCR page {page_no} failed: {e}")
            return ''
        logger.info(f"  OCR page {page_no}/{page_count}: {len(text)} chars")
        return text

    with ThreadPoolExecutor(max_workers=workers) as pool:
        texts = list(pool.map(ocr_page, range(1, page_count + 1)))

    return '\n'.join(t for t in texts if t.strip())


_OCR_CONFIG = r'--oem 3 --psm 6'
_PDFIUM_LOCK = threading.Lock()   # PDFium is not thread-safe

_TESSERACT_PATHS = [
    r'C:\Program Files\Tesseract-OCR\tesseract.exe',
    r'C:\Program Files (x86)\Tesseract-OCR\tesseract.exe',
    r'C:\Users\User\AppData\Local\Tesseract-OCR\tesseract.exe',
]
_POPPLER_SEARCH_PATHS = [
    r'C:\poppler\poppler-24.08.0\Library\bin',
    r'C:\Program Files\poppler\Library\bin',
    r'C:\Program Files\poppler-24.08.0\Library\bin',
    r'C:\poppler\Library\bin',
]


def _configure_tesseract_cmd(pytesseract) -> None:
    if shutil.which('tesseract'):
        return
    for tpath in _TESSERACT_PATHS:
        if os.path.exists(tpath):
            pytesseract.pytesseract.tesseract_cmd = tpath
            return


def _find_poppler_path():
    """
    Returns: '' when pdftoppm is on PATH (pdf2image needs no path), a
    directory when found in a known Windows location, or None when poppler
    is not installed at all.
    """
    if shutil.which('pdftoppm'):
        return ''
    for pp in _POPPLER_SEARCH_PATHS:
        if os.path.exists(os.path.join(pp, 'pdftoppm.exe')):
            return pp
    env = os.environ.get('POPPLER_PATH')
    if env and os.path.isdir(env):
        return env
    return None


def _pdfium_available() -> bool:
    try:
        import pypdfium2  # noqa: F401
        return True
    except ImportError:
        return False


def _page_count(filepath: str, poppler_path) -> int:
    if _pdfium_available():
        import pypdfium2 as pdfium
        with _PDFIUM_LOCK:
            doc = pdfium.PdfDocument(filepath)
            try:
                return len(doc)
            finally:
                doc.close()
    from pdf2image.pdf2image import pdfinfo_from_path
    return int(pdfinfo_from_path(filepath, poppler_path=poppler_path or None)['Pages'])


def _render_page_poppler(filepath: str, page_no: int, dpi: int, poppler_path):
    from pdf2image import convert_from_path
    images = convert_from_path(
        filepath, dpi=dpi, poppler_path=poppler_path or None, grayscale=True,
        first_page=page_no, last_page=page_no,
    )
    return images[0]


def _render_page_pdfium(filepath: str, page_no: int, dpi: int):
    import pypdfium2 as pdfium
    with _PDFIUM_LOCK:
        doc = pdfium.PdfDocument(filepath)
        try:
            page = doc[page_no - 1]
            try:
                bitmap = page.render(scale=dpi / 72, grayscale=True)
                # to_pil() shares the bitmap buffer — copy before it is released.
                return bitmap.to_pil().copy()
            finally:
                page.close()
        finally:
            doc.close()


def check_ocr_available() -> dict:
    """
    Check if OCR dependencies are available.
    Returns a dict with status info.
    """
    result = {
        'available': False,
        'pytesseract': False,
        'pdf2image': False,
        'pdfium': False,
        'tesseract_binary': False,
        'poppler_binary': False,
        'message': ''
    }

    try:
        import pytesseract  # noqa: F401
        result['pytesseract'] = True
    except ImportError:
        pass

    try:
        from pdf2image import convert_from_path  # noqa: F401
        result['pdf2image'] = True
    except ImportError:
        pass

    result['pdfium'] = _pdfium_available()

    # Tesseract binary: PATH first (Linux), then the known Windows installs
    if shutil.which('tesseract'):
        result['tesseract_binary'] = True
    else:
        for tpath in _TESSERACT_PATHS:
            if os.path.exists(tpath):
                result['tesseract_binary'] = True
                break

    # Poppler (pdftoppm/pdfinfo) — pdf2image shells out to it. On a fresh
    # Linux VPS this is the common silent failure (apt: poppler-utils).
    result['poppler_binary'] = _find_poppler_path() is not None

    rasteriser_ok = (result['pdf2image'] and result['poppler_binary']) or result['pdfium']
    result['available'] = all([result['pytesseract'], result['tesseract_binary'], rasteriser_ok])

    if result['available']:
        via = 'Poppler' if (result['pdf2image'] and result['poppler_binary']) else 'pypdfium2'
        result['message'] = f'OCR is fully available (Tesseract + {via})'
    else:
        missing = []
        if not result['pytesseract']:
            missing.append('pip install pytesseract')
        if not result['tesseract_binary']:
            missing.append('Install Tesseract (Linux: apt install tesseract-ocr)')
        if not rasteriser_ok:
            missing.append('Install Poppler (Linux: apt install poppler-utils) or pip install pypdfium2')
        result['message'] = 'Missing: ' + ', '.join(missing)

    return result
